from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


OUT = Path("docs/Phase3-Skills学习文档.docx")


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_style_font(style, name="宋体", size=10.5, bold=False):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=2):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, name="黑体", size=16 if level == 2 else 13.5, bold=True)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph(style="HTML Preformatted")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("• " + text)
    set_run_font(r)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
        shade_cell(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.175)
    sec.right_margin = Cm(3.175)
    set_style_font(doc.styles["Normal"])
    set_style_font(doc.styles["Heading 1"], "黑体", 24, True)
    set_style_font(doc.styles["Heading 2"], "黑体", 16, True)
    set_style_font(doc.styles["Heading 3"], "黑体", 13.5, True)
    if "HTML Preformatted" not in [style.name for style in doc.styles]:
        code_style = doc.styles.add_style("HTML Preformatted", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(doc.styles["HTML Preformatted"], "Consolas", 9)

    add_heading(doc, "Skills 功能学习文档", 2)
    add_para(doc, "本文针对当前 minimal-swe-agent 项目的 Extended M7 Skills 实现编写，目标是帮助面试时完整说明：Skill 是什么、为什么需要它、一次请求中如何发现和选择 Skill、正文如何进入模型上下文、外部 provider 如何连接、脚本如何安全执行，以及变更和失败如何恢复。本文以当前源码和测试为准，不把规划中的通用 MCP 平台或 Marketplace 生态能力写成已完成。")
    add_para(doc, "当前验证状态：npm run check 已通过，包含 102 个测试；typecheck、test 和 build 均通过。根据项目的里程碑约束，M4/M5 尚未按顺序完成正式发布验收，因此本文将其称为“Extended M7 工作区实现已验证”，而不是宣称 M7 正式发布完成。")

    add_heading(doc, "一、当前阶段实现概览", 2)
    add_para(doc, "当前 Skills 的设计不是简单把几个 Markdown 文件拼到 prompt 中，而是形成了一个有来源、有边界、有预算、有身份指纹、可刷新、可审计的上下文能力系统。Skill 的完整运行链路由本地发现、metadata 解析、资格判断、选择、正文读取、上下文注入和后续工具执行组成。")
    add_table(doc, ["能力", "当前状态", "主要入口"], [
        ("本地发现", "已实现：多 root、并发扫描、路径边界、错误隔离", "src/config/skills.ts"),
        ("选择与注入", "已实现：显式、lexical、embedding、hybrid 和 token 预算", "src/config/context-loader.ts"),
        ("增量刷新", "已实现：文件级 Skill 索引、watcher 定向失效", "src/skills/platform.ts"),
        ("脚本执行", "已实现：根目录约束、结构化 argv、ToolRouter 安全链", "src/config/skills.ts / src/skills/platform.ts"),
        ("MCP", "已实现：面向 Skills 的真实 stdio/HTTP client 和工具桥接", "src/skills/mcp-provider.ts"),
        ("远程 provider", "已实现：executor/orchestrator HTTP list/read/execute", "src/skills/remote-provider.ts"),
        ("插件市场", "已实现：搜索、semver、安装、升级、启停、卸载和依赖", "src/skills/plugin-manager.ts"),
        ("可观察性", "已实现：发现、读取、选择、安装、执行审计", "src/skills/platform.ts / src/index.ts"),
    ])
    add_para(doc, "面试时可以先给出一句总括：Skill 是模型上下文的可选专业能力包。系统常驻保存轻量 catalog，只有在显式或可靠选择后才注入正文；如果 Skill 需要操作文件、调用脚本或使用 MCP 工具，执行权限仍由通用工具安全链决定，选择 Skill 本身不会自动获得写文件或执行命令的权限。")

    add_heading(doc, "二、Skill 的数据模型与边界", 2)
    add_para(doc, "Skill 的最小身份由 id、name、description、path、canonicalPath、rootPath、scope、enabled、allowImplicitInvocation、body、bodyHash 和 diagnostics 组成。canonicalPath 是稳定身份，解决同名 Skill、符号链接和路径变化导致的身份漂移问题；bodyHash 用于判断正文是否变化，也用于上下文指纹和 Resume/Fork 后的重新验证。")
    add_para(doc, "Skill scope 当前分为 repo、user、system、admin。repo Skill 受项目 trust 影响，user/system/admin root 由受信配置声明。项目配置不能改变 user/system/admin root、Marketplace 地址、MCP server 或 enablement，这个限制防止项目通过自己的配置扩大能力边界。")
    add_para(doc, "上下文被分成三层：第一层是 catalog，只包含名称、描述和 locator；第二层是 selected body，选中后以独立的 user contextual fragment 注入完整 SKILL.md；第三层是 supporting resources，只有模型通过正常工具路径明确读取时才加载。这样既保留 Skill 的可发现性，又避免所有 Skill 正文常驻消耗上下文窗口。")

    add_heading(doc, "三、本地 Skill 发现流程", 2)
    add_para(doc, "本地发现的目标是从受控 root 中得到一个稳定的 SkillSnapshot，而不是尽可能多地读取文件。Snapshot 包含 skills、roots、fingerprint、catalogFingerprint 和 diagnostics，后续选择、上下文注入、watcher 刷新和会话恢复都使用这个快照。")
    add_heading(doc, "3.1 发现步骤", 3)
    for text in [
        "确定项目 root 和 trust。非受信项目不会读取 repo Skill 正文；不可信配置可以保留诊断，但不能改变 root 和权限推导。",
        "按 admin、system、user、repo 和兼容 root 建立候选 roots，并对路径做 resolve/realpath 校验。repo root 越出项目 root 时直接拒绝。",
        "对每个 root 执行有边界的递归扫描，限制最大深度、entry 数量和 Skill 数量；隐藏目录跳过，目录读取失败隔离为 warning。",
        "发现目录中的 SKILL.md 后解析 frontmatter，读取正文，计算 bodyHash，并读取可选的 agents/openai.yaml。",
        "以 canonicalPath 去重，以 name 检测同名冲突，应用 user/session enablement 规则，最后按 canonicalPath 稳定排序。",
        "计算 catalogFingerprint 和完整 fingerprint。catalogFingerprint 只覆盖可隐式选择的轻量 metadata，完整 fingerprint 还覆盖正文 hash、启用状态和诊断。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "3.2 SKILL.md 解析", 3)
    add_para(doc, "SKILL.md 必须以 YAML frontmatter 开始，description 必填，name 缺省时使用目录名，name 长度有限制，单行字段会规范化空白。disable-model-invocation 控制是否允许隐式选择，但它不代表执行授权。一个 Skill 解析失败只记录 invalid_skill 诊断，其他 Skill 仍然可以继续进入 catalog。")
    add_heading(doc, "3.3 openai.yaml 解析", 3)
    add_para(doc, "agents/openai.yaml 是可选的增强 metadata，当前支持 interface、policy.allow_implicit_invocation、products、dependencies.tools、MCP dependency 和 scripts。metadata 解析错误不会让有效 SKILL.md 消失，而是记录 invalid_skill_metadata；这样正文身份和可选配置的故障边界是分离的。")
    add_code(doc, "SKILL.md\n  -> frontmatter: name / description / policy\n  -> agents/openai.yaml: interface / products / dependencies / scripts\n  -> SkillMetadata { canonicalPath, bodyHash, diagnostics }")

    add_heading(doc, "四、并发扫描、缓存与文件级 watcher", 2)
    add_para(doc, "如果每次请求都递归读取所有 root，Skill 数量增加后会产生明显的 I/O 和 token 前置成本。因此当前实现将扫描拆成 root 任务并发执行，并缓存每个 root 的 snapshot 和“Skill 目录 -> metadata”索引。聚合 snapshot 仍然是不可变对象，避免 active request 被后台事件修改。")
    add_heading(doc, "4.1 缓存流程", 3)
    for text in [
        "以 cwd、roots、enablement、model 等配置生成 cache signature。",
        "在 TTL 内优先复用 root snapshot；不同 root 可以并行加载，scanConcurrency 控制并发上限。",
        "把每个 Skill 的父目录记录到 root entries Map，作为增量刷新时的反向定位索引。",
        "聚合各 root 的 skills、roots 和 diagnostics，重新生成排序稳定的 SkillSnapshot。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "4.2 文件级 watcher 流程", 3)
    add_para(doc, "watcher 接收 root、事件类型和 filename。已知 Skill 目录中的 SKILL.md 或 agents/openai.yaml 变化时，只重新解析该目录；删除 SKILL.md 时删除该目录索引项；新增 Skill 目录或 filename 缺失、rename/overflow 无法判断时，才回退该 root 重扫。增量更新完成后，聚合缓存失效，下一次 context preparation 才生成新上下文。")
    add_code(doc, "fs.watch(root)\n  -> changedPath\n  -> locate Skill directory\n  -> loadSkillDirectory(one directory)\n  -> replace/remove one index entry\n  -> recompute fingerprints\n  -> next context refresh uses new immutable snapshot")
    add_para(doc, "这套设计的关键取舍是：索引更新可以增量化，但请求快照不能原地修改。这样 watcher 不会在模型请求中途改变已经发送的 prompt，也不会破坏 Resume/Fork 的历史一致性。")

    add_heading(doc, "五、Skill 选择与上下文注入", 2)
    add_heading(doc, "5.1 显式选择", 3)
    add_para(doc, "显式选择是最稳定的路径。turn-runner 从用户请求中提取 $skill-name，然后 context-loader 解析 name、绝对路径、canonical path 或 provider locator。名称不唯一时返回 skill_name_ambiguous，要求调用方使用更精确的 locator，不能按排序结果猜一个。")
    add_heading(doc, "5.2 隐式选择", 3)
    add_para(doc, "只有开启 implicit_selection 且 Skill 的 allowImplicitInvocation 为 true 时才进入隐式选择。lexical 模式按 query terms 与 name/description 的重合度计算；embedding 模式调用 OpenAI-compatible embeddings endpoint；hybrid 模式将 semantic score 和 lexical score 加权合并。结果按 score 和 canonical identity 稳定排序，两个候选得分过近时返回 ambiguous 并跳过自动注入。")
    add_para(doc, "embedding client 具备批处理、超时、取消、有限重试和向量维度校验。候选向量缓存的 key 包含模型、Skill identity、bodyHash 和 catalog text，因此正文或描述改变后不会错误复用旧向量。embedding endpoint 没有配置或请求失败时，系统生成结构化诊断，不把 lexical fallback 冒充成语义选择成功。")
    add_heading(doc, "5.3 资格校验", 3)
    add_para(doc, "选择前会执行 products 和 dependencies 门禁。Skill 声明的 product 必须与当前 products 集合匹配；tool dependency 必须存在于 ToolRegistry；MCP dependency 必须对应就绪 provider。缺少依赖的 Skill 会被标记为 disabled 并给出 warning，隐式选择会跳过，显式选择则得到可解释错误。")
    add_heading(doc, "5.4 注入流程", 3)
    for text in [
        "turn-runner 为每个普通 turn 提取显式 mention，并把用户请求作为隐式选择 query。",
        "context-loader 加载 AGENTS snapshot 和 local/provider Skill snapshot。",
        "执行产品与依赖资格校验，然后确定 effectiveMentions。",
        "按 catalog budget 渲染 developer fragment；按 selected body budget 构造 user fragment。",
        "远程 provider Skill 直到被选中才执行 read，并把返回正文 hash 化后重新构建 SkillContext。",
        "将 instructions、catalog 和 selected body 写入 AgentContext，同时更新 world/reference/compHash 身份。",
    ]:
        add_bullet(doc, text)
    add_code(doc, "user request\n  -> extract $mentions\n  -> refresh AGENTS + Skill snapshots\n  -> eligibility check\n  -> explicit / lexical / embedding / hybrid selection\n  -> catalog fragment + selected <skill> body fragment\n  -> model request")

    add_heading(doc, "六、Token 预算与 tokenizer 选型", 2)
    add_para(doc, "Skill 上下文至少有三类成本：catalog metadata、selected Skill 正文和模型请求中的其他历史/工具内容。正文支持单 Skill 上限和单轮总上限，超限时拒绝选择，而不是无提示截断指令。")
    add_para(doc, "当前主路径使用 js-tiktoken。选择它的原因是它提供真实的 BPE 编码实现，能对已知 OpenAI 模型使用对应 encoding；相比按字符或中英文比例估算，更适合做硬预算和面试时解释 token accounting。对于未知 provider model，系统使用明确声明的 fallback encoding，并通过 describe() 标记 modelMatched=false，不把 fallback 伪装成精确模型 tokenizer。")
    add_table(doc, ["预算对象", "处理方式", "超限行为"], [
        ("catalog", "只放 name/description/locator，按预算截取描述", "省略尾部或缩短描述并生成诊断"),
        ("selected body", "使用真实 tokenizer 计算正文 token", "拒绝选择，不静默破坏正文语义"),
        ("remote read", "读取后计算 bodyHash，再重新构建上下文", "读取失败则诊断并不注入正文"),
    ])

    add_heading(doc, "七、Skill 脚本执行安全链", 2)
    add_para(doc, "Skill 脚本的作用是让 Skill 提供可重复的局部自动化能力，例如生成报告、运行格式化器或调用 Skill 自带的小工具。脚本不能因为放在 Skill 目录里就被视为可信命令，当前实现把它转换为普通 ToolRegistration，让它复用现有工具安全模型。")
    add_heading(doc, "7.1 脚本解析", 3)
    for text in [
        "openai.yaml 中脚本必须声明 name、description 和 path，可选 cwd、args、parameters、read_only。",
        "path 必须位于 Skill root 内，解析后必须是实际存在的 .js/.mjs/.cjs 文件。",
        "cwd 必须位于 Skill root 内，并通过 realpath 检查符号链接逃逸。",
        "参数禁止包含换行和 NUL 等控制字符，避免把 metadata 变成额外命令语句。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "7.2 执行流程", 3)
    add_para(doc, "脚本 metadata 先转换为 ToolSpec 和 ToolRuntime。模型产生 tool call 后，ToolRouter 执行参数 schema 校验、路径 preflight、风险分析、preview、PermissionPolicy、ApprovalBroker、审批后 revalidation，然后调用 sandboxProvider。实际启动使用结构化 argv，避免 Windows cmd 或 Unix shell 对绝对路径和引号的二次解释；输入通过受限环境变量传递，stdout/stderr 和 executionId 记录为 ToolResult。")
    add_code(doc, "model tool_call\n  -> ToolRegistry schema validation\n  -> ToolPreflight / workspace policy\n  -> preview + permission\n  -> approval\n  -> revalidation\n  -> SandboxProvider(argv, fixed cwd, filtered env)\n  -> ToolResult + Audit")

    add_heading(doc, "八、MCP 与远程 Skill provider", 2)
    add_para(doc, "当前阶段的 MCP 定位是为 Skills 和工具提供外部能力，不是完整的通用 MCP 平台管理器。实现使用 @modelcontextprotocol/sdk，支持 stdio 和 Streamable HTTP 两类 transport。每个 provider 拥有自己的 client、transport 和连接状态，某个 MCP 服务失败不会直接终止整个 agent session。")
    add_heading(doc, "8.1 MCP Skill 流程", 3)
    for text in [
        "从受信配置创建 McpSkillProvider，stdio 使用受限环境启动命令，HTTP 使用 StreamableHTTPClientTransport。",
        "首次 list/read/execute 时建立 MCP handshake，并读取 server capabilities。",
        "如果支持 resources，则 listResources 将 resource 映射为 SkillMetadata；resource URI 通过 resourcePrefixes 过滤。",
        "选中远程 resource 后调用 readResource，把 text 内容作为 Skill body，并计算 bodyHash。",
        "如果支持 tools，则读取工具 schema，注册为 mcp.<server>.<tool>，调用仍经过 ToolRouter 的权限和审计链。",
        "请求有 timeout、AbortSignal cancellation 和 degraded/closed 状态；close 会关闭 client 和 transport。",
    ]:
        add_bullet(doc, text)
    add_para(doc, "executor/orchestrator provider 使用项目自己的 HTTP provider 协议，提供 list、read、execute 三个操作，同样具备认证、超时、取消、响应校验和错误隔离。它们是统一 SkillProvider 的远程实现，不代表已经兼容所有外部平台的协议。")

    add_heading(doc, "九、Plugin 与 Marketplace 生命周期", 2)
    add_para(doc, "插件的目标是把一组 Skill、工具和 provider 配置作为可版本化单元管理。当前生命周期由 PluginLifecycleManager 负责，MarketplaceClient 负责 manifest-first 获取和校验，MarketplaceIndex 负责索引搜索与 semver 解析。")
    add_heading(doc, "9.1 安装流程", 3)
    for text in [
        "CLI 接收 skills install <url>，记录 preflight audit。",
        "MarketplaceClient 获取 manifest/files，检查 HTTP 状态、manifest schema、contentHash、trust policy 和授权回调。",
        "PluginLifecycleManager 校验 product、requiredPermissions 和依赖版本；发现循环依赖或缺少版本时失败。",
        "文件写入 staging 目录，逐个检查相对路径和 symlink，写入 plugin.json 后执行目录替换。",
        "替换失败时恢复 backup，成功后更新 registry.json，记录 activeVersion、versions 和 enabled。",
        "Skill root 只暴露已启用插件的活动版本，旧版本不会自动进入 catalog。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "9.2 运维命令", 3)
    add_code(doc, "node dist/index.js skills list\nnode dist/index.js skills search <query>\nnode dist/index.js skills install <url>\nnode dist/index.js skills upgrade <id> [range]\nnode dist/index.js skills enable <id>\nnode dist/index.js skills disable <id>\nnode dist/index.js skills uninstall <id> [version]")
    add_para(doc, "当前边界是：项目提供 Marketplace 协议、hash/trust 回调和生命周期，但不内置第三方签名信任根、账号体系、支付体系或 Marketplace SLA。这些是部署方和具体市场服务的职责。")

    add_heading(doc, "十、审计、失败处理与会话一致性", 2)
    add_para(doc, "Skills 不能绕过项目已有安全规则。发现、读取、选择、安装和执行都在关键入口生成 audit；记录 source、operation、hash、bytes、decision 和 success，但不记录完整正文、API key、完整环境变量或凭据。")
    add_heading(doc, "10.1 失败分类", 3)
    add_table(doc, ["失败类型", "处理方式", "是否阻止整个会话"], [
        ("单个 SKILL.md 损坏", "记录 invalid_skill，隔离该项", "否"),
        ("openai.yaml 错误", "保留正文，忽略可选 metadata 并告警", "否"),
        ("embedding 服务失败", "记录结构化诊断，不伪装成语义成功", "通常否；显式要求时报告错误"),
        ("MCP provider 断线", "provider degraded，其他 provider 继续", "否"),
        ("脚本审批拒绝或 sandbox 能力不足", "ToolResult isError，保留审计", "否"),
        ("Marketplace hash/trust/依赖失败", "安装不提交，staging 清理或恢复 backup", "只阻止当前安装"),
    ])
    add_heading(doc, "10.2 与 Resume/Fork/Compaction 的关系", 3)
    add_para(doc, "每轮 context refresh 都重新发现当前 AGENTS 和 Skills，但历史中保存的是当轮 snapshot identity、selected Skill locator 和 bodyHash。Resume/Fork 先重放历史快照，再在下一普通 turn 比较当前文件系统指纹；如果发生变化，系统注入 replacement/removal 或当前 full context，而不是把新正文伪装成旧历史。compaction 只压缩历史，不把 Skill catalog 和 selected body 的身份丢掉。")

    add_heading(doc, "十一、技术选型与原因", 2)
    add_table(doc, ["技术", "使用位置", "选择原因"], [
        ("TypeScript strict", "Skill metadata、provider、CLI 和运行时", "利用现有工程类型边界，减少 provider 和安全链的隐式状态错误"),
        ("yaml", "SKILL.md frontmatter、openai.yaml", "当前项目已有 YAML 依赖，适合解析嵌套 metadata 并支持唯一键校验"),
        ("js-tiktoken", "正文和 catalog token 预算", "提供真实 BPE 计数，避免字符估算与真实模型窗口偏差"),
        ("@modelcontextprotocol/sdk", "MCP stdio/HTTP client", "直接实现 MCP handshake、resource、tool 和 capability 协议，避免手写 JSON-RPC 客户端"),
        ("semver", "插件版本选择和依赖", "标准化范围匹配、最高兼容版本和版本排序"),
        ("Node fs.watch", "Skill root watcher", "跨平台可用的基础事件源；上层用索引、路径判定和 root fallback 补足不确定事件"),
        ("SHA-256", "body、catalog、snapshot、embedding cache identity", "稳定、可复现、适合做内容身份和缓存键，不把路径或时间戳当正文身份"),
    ])
    add_para(doc, "没有为 Skill 单独发明另一套权限系统，原因是当前项目已有 ToolRouter、WorkspacePolicy、PermissionPolicy、ApprovalBroker、SandboxProvider 和 AuditTrail。脚本和 MCP tool 复用这条链，能保证“上下文能力”和“执行授权”保持分离。")

    add_heading(doc, "十二、面试时如何完整讲解一次请求", 2)
    add_para(doc, "可以按下面的顺序回答，既能说明功能，也能说明工程约束：")
    for text in [
        "用户输入“使用 $database skill 分析迁移风险”。",
        "turn-runner 提取 database mention，context-loader 在受控 roots 中刷新 AGENTS 和 Skills snapshot。",
        "本地 provider 使用缓存；如果 watcher 已更新 database 目录，则只解析该目录并更新 bodyHash。",
        "系统检查 database Skill 是否启用、是否允许显式调用、是否满足 products 和 tool/MCP dependencies。",
        "selector 用 canonical locator 精确找到 Skill，构造 catalog fragment 和 selected body fragment，并用 tiktoken 检查预算。",
        "world state/reference context 记录 catalogFingerprint、selectedSkillFingerprint、bodyHash 和 compHash。",
        "模型收到包含 Skill 正文的请求，产生 read_file 或 MCP tool call。",
        "工具调用进入 ToolRouter，经历 schema、preflight、preview、permission、approval、revalidation、sandbox、audit。",
        "如果用户修改了 Skill 文件，watcher 只更新该 Skill 索引；当前请求不变，下一轮重新注入新版本。",
        "Resume/Fork 后使用历史 snapshot identity 判断上下文是否一致，不把当前文件内容错误地当作历史内容。",
    ]:
        add_bullet(doc, text)
    add_code(doc, "请求 -> 发现 -> 资格 -> 选择 -> 预算 -> 注入 -> 模型 tool call\n     -> ToolRouter 安全执行 -> 审计 -> snapshot/replay 一致性")

    add_heading(doc, "十三、当前边界与后续方向", 2)
    add_para(doc, "当前核心 Skills 功能已经形成真实闭环，但仍有明确边界：Marketplace 没有内置第三方签名信任根；MCP 还需要后续补充分页、服务变更订阅、健康检查、自动重连和更完整 OAuth 生命周期；watcher 在无法定位事件时会回退 root 重扫，尚不是跨进程文件级索引服务；embedding 依赖外部 embedding endpoint；executor/orchestrator 使用项目定义的 HTTP 协议。")
    add_para(doc, "这些边界不影响当前面试项目展示的核心价值，因为本项目已经证明了发现、选择、预算、provider、执行、安全、持久化和测试之间的完整关系；但在介绍时必须把“已实现的主路径”和“需要部署能力或后续阶段的生态能力”分开。")

    add_heading(doc, "十四、总结", 2)
    add_para(doc, "当前 Skills 实现的核心不是读取 SKILL.md，而是建立了一个可解释的上下文能力系统：所有输入都有来源和 canonical identity，发现有 root 和 trust 边界，选择有显式与语义路径，正文有真实 token 预算，provider 有生命周期和故障隔离，脚本和 MCP tool 不能绕过统一安全链，文件变化可以通过增量索引在下一轮刷新生效，Resume/Fork/Compaction 仍然能够验证上下文身份。")
    add_para(doc, "对于求职和面试，这个阶段最值得展示的工程判断有三点。第一，Skill 上下文和执行授权被明确分离，避免“注入了指令就能执行命令”的危险设计。第二，缓存和 watcher 不是简单清空缓存，而是用文件级索引降低不必要的扫描，同时保留不确定事件的安全回退。第三，外部能力全部经过 provider、ToolRouter、sandbox 和 audit 边界，失败时有诊断和恢复语义，而不是用 mock 或接口存在来假装能力已经完成。")
    add_para(doc, "最终验证：npm run typecheck、npm test、npm run build 均通过，当前测试总数为 102。")

    # Keep the document close to the supplied note style: no cover page, broad margins,
    # numbered section headings, and compact explanatory paragraphs.
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
