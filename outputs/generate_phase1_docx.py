from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src = Path(r"C:\Users\LENOVO\AppData\Roaming\WorkBuddy\User\globalStorage\tencent-cloud.coding-copilot\brain\97dca546e9ec4b0ca84624bf75859654\Phase1学习笔记.md")
out = src.with_suffix(".docx")
text = src.read_text(encoding="utf-8")
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

for style_name, size, bold in [("Normal", 10.5, False), ("Title", 22, True), ("Heading 1", 16, True), ("Heading 2", 13, True), ("Heading 3", 11.5, True)]:
    style = doc.styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(rows):
    data = []
    for row in rows:
        if re.fullmatch(r"\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*", row):
            continue
        data.append([item.strip() for item in row.strip().strip("|").split("|")])
    if not data:
        return
    columns = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(data):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = row[col_index] if col_index < len(row) else ""
            if row_index == 0:
                shade(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(9)
    doc.add_paragraph()


lines = text.splitlines()
in_code = False
code_lines = []
table_lines = []
for line in lines:
    if line.startswith("```"):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            in_code = False
        continue
    if in_code:
        code_lines.append(line)
        continue
    if line.startswith("|"):
        table_lines.append(line)
        continue
    if table_lines:
        add_table(table_lines)
        table_lines = []
    if not line.strip() or line.strip() == "---":
        continue
    if line.startswith("# "):
        paragraph = doc.add_paragraph(style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(line[2:].strip())
    elif line.startswith("### "):
        doc.add_paragraph(line[4:].strip(), style="Heading 3")
    elif line.startswith("## "):
        doc.add_paragraph(line[3:].strip(), style="Heading 1")
    elif re.match(r"^\d+\.\s+", line):
        doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
    elif re.match(r"^[-*]\s+", line):
        doc.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
    else:
        doc.add_paragraph(line)
if table_lines:
    add_table(table_lines)

doc.core_properties.title = "Phase 1 学习笔记"
doc.core_properties.subject = "minimal-swe-agent Phrase 1 技术学习笔记"
doc.save(out)
print(out)
